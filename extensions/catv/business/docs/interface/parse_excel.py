#!/usr/bin/env python3
"""
엑셀 파일을 interface_data.json으로 변환하는 스크립트
모든 시트의 셀 데이터를 상세하게 파싱합니다.
"""

import json
import os
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.cell.cell import MergedCell


def parse_excel_to_json(excel_path, output_path):
    """엑셀 파일을 JSON으로 변환"""

    print(f"엑셀 파일 로딩 중: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    data = {
        'file_info': {
            'source_file': os.path.basename(excel_path),
            'parsed_at': datetime.now().isoformat(),
            'total_sheets': len(wb.sheetnames)
        },
        'sheets': []
    }

    for sheet_name in wb.sheetnames:
        print(f"  시트 파싱 중: {sheet_name}")
        ws = wb[sheet_name]

        sheet_data = {'name': sheet_name, 'rows': []}

        # 시트의 모든 행 처리
        for row_idx, row in enumerate(ws.iter_rows()):
            row_data = {'row_index': row_idx, 'cells': []}

            for col_idx, cell in enumerate(row):
                cell_data = {
                    'col_index': col_idx,
                    'value': None,
                    'merged': False
                }

                # 셀 값 가져오기
                if cell.value is not None:
                    if isinstance(cell.value, (int, float)):
                        cell_data['value'] = cell.value
                    else:
                        cell_data['value'] = str(cell.value)

                # 병합 셀 확인
                if isinstance(cell, MergedCell):
                    cell_data['merged'] = True

                row_data['cells'].append(cell_data)

            sheet_data['rows'].append(row_data)

        data['sheets'].append(sheet_data)

    # JSON 저장
    print(f"\nJSON 파일 저장 중: {output_path}")
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    # 통계 출력
    total_rows = sum(len(sheet['rows']) for sheet in data['sheets'])
    print(f"\n✅ 변환 완료!")
    print(f"  시트 수: {len(data['sheets'])}")
    print(f"  총 행 수: {total_rows}")
    print(f"  파일 크기: {os.path.getsize(output_path):,} bytes")


def main():
    """메인 함수"""
    import sys

    # 현재 스크립트 디렉토리
    script_dir = Path(__file__).parent

    # 명령줄 인자로 파일 경로 받기
    if len(sys.argv) > 1:
        excel_path = Path(sys.argv[1])
        if not excel_path.is_absolute():
            excel_path = script_dir / excel_path
    else:
        # 기본값: reference 폴더에서 최신 버전 찾기
        reference_dir = script_dir / 'reference'
        excel_files = sorted(reference_dir.glob('*v2.*.xlsx'), reverse=True)

        if not excel_files:
            print("❌ 엑셀 파일을 찾을 수 없습니다.")
            print(f"   찾은 위치: {reference_dir}")
            print("\n사용법: python3 parse_excel.py [엑셀파일경로]")
            return 1

        excel_path = excel_files[0]

    if not excel_path.exists():
        print(f"❌ 파일을 찾을 수 없습니다: {excel_path}")
        return 1

    output_path = script_dir / 'interface_data.json'

    print("=" * 70)
    print("엑셀 → JSON 변환 스크립트")
    print("=" * 70)
    print(f"입력: {excel_path.name}")
    print(f"출력: {output_path.name}")
    print("=" * 70)
    print()

    try:
        parse_excel_to_json(excel_path, output_path)
        return 0
    except Exception as e:
        print(f"\n❌ 오류 발생: {e}")
        import traceback
        traceback.print_exc()
        return 1


if __name__ == '__main__':
    exit(main())
